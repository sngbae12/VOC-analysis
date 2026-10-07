# -*- coding: utf-8 -*-
"""
11장 실습: 고객사 VOC 분석 Agent (CrewAI + Gradio)

구성
  - 로컬 처리    : CSV 로드, 통계/막대그래프, 워드클라우드, DOCX 저장
  - Issue Agent : 보고서 이슈 문장만 OpenAI/CrewAI로 생성

실행
  python voc_analysis_app.py

패키지(로컬 환경 버전 우선)
  crewai 1.15.1, gradio 6.19.0, plotly 6.8.0, pandas 3.0.3,
  wordcloud 1.9.6, python-docx 1.2.0, openai 2.44.0, kaleido 1.3.0
"""

import hashlib
import html
import io
import logging
import os
import queue
import re
import socket
import sys
import threading
import traceback
from collections import Counter
from datetime import datetime
from pathlib import Path
from typing import Any, Callable, Iterator

os.environ.setdefault("CREWAI_DISABLE_TELEMETRY", "true")
os.environ.setdefault("OTEL_SDK_DISABLED", "true")

import matplotlib

matplotlib.use("Agg")
import matplotlib.pyplot as plt
from matplotlib import font_manager
import pandas as pd
import plotly.graph_objects as go
from dotenv import load_dotenv
from wordcloud import WordCloud

load_dotenv(Path(__file__).resolve().parent / ".env")

from crewai import Agent, Crew, LLM, Process, Task
from crewai.tools import tool
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
import gradio as gr


# ---------------------------------------------------------------------------
# 경로·상수
# ---------------------------------------------------------------------------
BASE_DIR = Path(__file__).resolve().parent
OUTPUT_DIR = BASE_DIR / "output"
EXPORT_DIR = OUTPUT_DIR / "exports"
UPLOAD_DIR = OUTPUT_DIR / "uploads"
SAMPLE_CSV = BASE_DIR / "sample_voc.csv"
for _dir in (OUTPUT_DIR, EXPORT_DIR, UPLOAD_DIR):
    _dir.mkdir(parents=True, exist_ok=True)

REQUIRED_COLUMNS = ["순번", "일자", "고객명", "산업군", "지역", "제품명", "분야", "불만"]
STAT_COLUMNS = ["산업군", "제품명", "분야"]

MODEL_NAME = (
    os.getenv("OPENAI_MODEL_NAME")
    or os.getenv("MODEL")
    or "gpt-4o-mini"
)

KOREAN_STOPWORDS = {
    "그리고", "그러나", "하지만", "또는", "및", "등", "관련", "대한", "통해", "위해",
    "에서", "으로", "로서", "로서의", "하다", "되다", "있다", "없다", "이다", "아니다",
    "한다", "됩니다", "입니다", "했습니다", "합니다", "했습니다만", "요청", "문의",
    "사항", "부분", "경우", "때문", "때문에", "너무", "매우", "조금", "계속", "이번",
    "해당", "현재", "이후", "이전", "오늘", "내일", "저희", "우리", "고객", "사용",
    "문제", "발생", "정도", "가능", "필요", "확인", "처리", "진행", "관련해",
    "있습니다", "없습니다", "됩니다", "합니다", "했습니다", "되었습니다", "됩니다만",
    "같습니다", "됩니다", "합니다만", "합니다", "해서", "되어", "되는", "하는",
    "있는", "없는", "같은", "이런", "그런", "어떤", "여러", "모든", "일부", "각각",
    "자주", "계속", "반복", "간헐적", "간헐적으로", "어렵습니다", "안됩니다",
}

NAVY = RGBColor(0x1F, 0x3A, 0x5F)
ACCENT = RGBColor(0x2F, 0x6F, 0xB3)
DARK = RGBColor(0x22, 0x22, 0x22)


_EAST_ASIA_NAMES = {
    "Malgun Gothic": "맑은 고딕",
    "Gulim": "굴림",
    "Batang": "바탕",
    "Hancom Gothic": "한컴고딕",
}


def _register_font_file(font_path: Path) -> tuple[str, str, str] | None:
    if not font_path.exists():
        return None
    try:
        font_manager.fontManager.addfont(str(font_path))
    except Exception:
        pass
    try:
        prop = font_manager.FontProperties(fname=str(font_path))
        family = (prop.get_name() or "").strip()
    except Exception:
        return None
    if not family:
        return None
    east = _EAST_ASIA_NAMES.get(family, family)
    return str(font_path), family, east


def _find_korean_font() -> tuple[str, str, str]:
    """(ttf 경로, Plotly/Matplotlib 패밀리명, Word 동아시아 폰트명)을 반환합니다."""
    candidates: list[Path] = []
    custom = (os.getenv("VOC_FONT_PATH") or "").strip()
    if custom:
        candidates.append(Path(custom).expanduser())
    fonts_dir = Path(os.environ.get("WINDIR", r"C:\Windows")) / "Fonts"
    candidates.extend(
        [
            fonts_dir / "malgun.ttf",
            fonts_dir / "malgunbd.ttf",
            fonts_dir / "Hancom Gothic Regular.ttf",
            fonts_dir / "gulim.ttc",
            fonts_dir / "batang.ttc",
        ]
    )
    linux_fonts = [
        Path("/usr/share/fonts/truetype/nanum/NanumGothic.ttf"),
        Path("/usr/share/fonts/truetype/nanum/NanumBarunGothic.ttf"),
        Path("/usr/share/fonts/truetype/nanum/NanumGothicBold.ttf"),
        Path("/usr/share/fonts/opentype/nanum/NanumGothic.ttf"),
        Path("/usr/share/fonts/truetype/noto/NotoSansCJK-Regular.ttc"),
        Path("/usr/share/fonts/opentype/noto/NotoSansCJK-Regular.ttc"),
        Path("/usr/share/fonts/opentype/noto/NotoSansCJKkr-Regular.otf"),
        Path("/usr/share/fonts/truetype/google-noto-cjk/NotoSansCJK-Regular.ttc"),
    ]
    candidates.extend(linux_fonts)
    for font_path in candidates:
        found = _register_font_file(font_path)
        if found:
            return found
    return "", "sans-serif", "Malgun Gothic"


FONT_PATH, FONT_FAMILY, FONT_EAST_ASIA = _find_korean_font()
FONT_WARNING = ""
if FONT_PATH:
    plt.rcParams["font.family"] = FONT_FAMILY
else:
    FONT_WARNING = (
        "한글 폰트(맑은 고딕, NanumGothic 등)를 찾지 못했습니다. "
        "Windows 폰트 폴더를 확인하거나 VOC_FONT_PATH에 ttf 경로를 지정하세요."
    )
plt.rcParams["axes.unicode_minus"] = False


STORE: dict[str, Any] = {
    "df": None,
    "csv_path": None,
    "source_kind": None,
    "source_label": None,
    "charts": {},
    "stats": {},
    "wordcloud_path": None,
    "keywords": {},
    "issue_analysis": "",
    "report_path": None,
    "chart_errors": {},
}
STORE_LOCK = threading.Lock()
_RUNTIME = threading.local()

SUPPORTED_PYTHON = ((3, 11), (3, 12), (3, 13))
RECOMMENDED_PYTHON = (3, 11)
APP_PORT = 7860
DOWNLOAD_EXPORT_SUFFIXES = {".png", ".docx"}
_EXPORT_DIGESTS: set[str] = set()
_EXPORT_LOCK = threading.Lock()
_GUARD_INSTALLED = False
_EMPTY_DATE_TOKENS = {"", "nan", "none", "nat", "<na>", "null", "na"}


def python_support_message(info: Any = None) -> tuple[bool, str]:
    info = info or sys.version_info
    major_minor = (int(info.major), int(info.minor))
    rec = ".".join(str(x) for x in RECOMMENDED_PYTHON)
    if major_minor in SUPPORTED_PYTHON:
        return True, ""
    current = f"{info.major}.{info.minor}.{getattr(info, 'micro', 0)}"
    return False, (
        f"This app supports Python 3.11, 3.12, and 3.13 only. "
        f"Current version is Python {current}. "
        f"Pinned CrewAI 1.15.1 does not allow Python 3.14, and pandas 3.0 needs 3.11 or newer. "
        f"Python {rec} is recommended."
    )


def port_in_use(port: int, host: str = "127.0.0.1") -> bool:
    with socket.socket(socket.AF_INET, socket.SOCK_STREAM) as sock:
        sock.settimeout(0.4)
        return sock.connect_ex((host, port)) == 0


def default_blocked_paths() -> list[str]:
    return [
        str((BASE_DIR / ".env").resolve()),
        str((BASE_DIR / ".git").resolve()),
        str((BASE_DIR / ".venv").resolve()),
        str((BASE_DIR / "venv").resolve()),
        str((BASE_DIR / ".verify_venv").resolve()),
        str(UPLOAD_DIR.resolve()),
        str((BASE_DIR / "voc_analysis_app.py").resolve()),
        str(SAMPLE_CSV.resolve()),
        str((BASE_DIR / "requirements.txt").resolve()),
        str((BASE_DIR / ".env.example").resolve()),
    ]


def file_sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with Path(path).open("rb") as handle:
        for chunk in iter(lambda: handle.read(65536), b""):
            digest.update(chunk)
    return digest.hexdigest()


def register_export_file(path: Path | None) -> Path | None:
    if path is None:
        return None
    candidate = Path(path)
    if not candidate.exists() or not candidate.is_file():
        return None
    if candidate.suffix.lower() not in DOWNLOAD_EXPORT_SUFFIXES:
        return None
    digest = file_sha256(candidate)
    with _EXPORT_LOCK:
        _EXPORT_DIGESTS.add(digest)
    return candidate


def _is_blocked_location(resolved: Path) -> bool:
    env_file = (BASE_DIR / ".env").resolve()
    if resolved == env_file or resolved.name.lower().startswith(".env"):
        return True
    for root in (
        BASE_DIR / ".git",
        BASE_DIR / ".venv",
        BASE_DIR / "venv",
        BASE_DIR / ".verify_venv",
        UPLOAD_DIR,
    ):
        try:
            resolved.relative_to(root.resolve())
            return True
        except (ValueError, OSError):
            continue
    if resolved.suffix.lower() in {".py", ".pyc", ".bat", ".csv", ".md", ".json", ".toml"}:
        return True
    return False


def is_safe_download_path(path: Path) -> bool:
    try:
        resolved = Path(path).resolve()
    except OSError:
        return False
    if not resolved.exists() or not resolved.is_file():
        return False
    if _is_blocked_location(resolved):
        return False
    if resolved.suffix.lower() not in DOWNLOAD_EXPORT_SUFFIXES:
        return False
    try:
        digest = file_sha256(resolved)
    except OSError:
        return False
    with _EXPORT_LOCK:
        return digest in _EXPORT_DIGESTS


def install_download_guard() -> None:
    global _GUARD_INSTALLED
    if _GUARD_INSTALLED:
        return
    from fastapi import HTTPException
    from gradio import route_utils, routes as gr_routes, utils as gr_utils
    from gradio import static_server as gr_static
    from gradio_client import utils as client_utils

    original = route_utils.file_fetch

    def guarded_file_fetch(path_or_url, request, blocks_or_config, upload_dir):
        if client_utils.is_http_url_like(path_or_url):
            return original(path_or_url, request, blocks_or_config, upload_dir)
        try:
            abs_path = Path(gr_utils.abspath(path_or_url))
        except Exception as exc:
            raise HTTPException(403, f"File not allowed: {path_or_url}.") from exc
        if abs_path.exists() and abs_path.is_file() and not is_safe_download_path(abs_path):
            raise HTTPException(403, f"File not allowed: {path_or_url}.")
        return original(path_or_url, request, blocks_or_config, upload_dir)

    route_utils.file_fetch = guarded_file_fetch
    gr_routes.file_fetch = guarded_file_fetch
    gr_static.file_fetch = guarded_file_fetch
    _GUARD_INSTALLED = True


def _date_token(value: Any) -> str:
    if value is None:
        return ""
    try:
        if pd.isna(value):
            return ""
    except (TypeError, ValueError):
        pass
    if isinstance(value, pd.Timestamp):
        if pd.isna(value):
            return ""
        return value.strftime("%Y-%m-%d")
    if isinstance(value, datetime):
        return value.strftime("%Y-%m-%d")
    if isinstance(value, bool):
        return str(value)
    if isinstance(value, int):
        if 10000101 <= value <= 29991231:
            return f"{value:08d}"
        return str(value)
    if isinstance(value, float):
        if value.is_integer() and 10000101 <= int(value) <= 29991231:
            return f"{int(value):08d}"
        return str(value)
    text = str(value).strip()
    if re.fullmatch(r"\d{8}(?:\.0+)?", text):
        return text.split(".", 1)[0]
    return text


def parse_voc_dates(series: pd.Series) -> tuple[pd.Series, int, int]:
    """Parse YYYYMMDD integers/strings and ISO dates without treating ints as ns."""
    result = pd.Series(pd.NaT, index=series.index, dtype="datetime64[ns]")
    if pd.api.types.is_datetime64_any_dtype(series):
        parsed = pd.to_datetime(series, errors="coerce")
        n_invalid = int(parsed.isna().sum())
        return parsed, n_invalid, 0
    tokens = series.map(_date_token)
    empty_mask = tokens.str.lower().isin(_EMPTY_DATE_TOKENS)
    remaining = ~empty_mask

    ymd_mask = remaining & tokens.str.fullmatch(r"\d{8}", na=False)
    if ymd_mask.any():
        result.loc[ymd_mask] = pd.to_datetime(tokens.loc[ymd_mask], format="%Y%m%d", errors="coerce")
        remaining = remaining & result.isna()

    iso_mask = remaining & tokens.str.fullmatch(r"\d{4}-\d{1,2}-\d{1,2}", na=False)
    if iso_mask.any():
        result.loc[iso_mask] = pd.to_datetime(tokens.loc[iso_mask], errors="coerce")
        remaining = remaining & result.isna()

    slash_mask = remaining & tokens.str.fullmatch(r"\d{4}/\d{1,2}/\d{1,2}", na=False)
    if slash_mask.any():
        result.loc[slash_mask] = pd.to_datetime(tokens.loc[slash_mask], errors="coerce")
        remaining = remaining & result.isna()

    dot_mask = remaining & tokens.str.fullmatch(r"\d{4}\.\d{1,2}\.\d{1,2}", na=False)
    if dot_mask.any():
        result.loc[dot_mask] = pd.to_datetime(tokens.loc[dot_mask], errors="coerce")
        remaining = remaining & result.isna()

    still = remaining & result.isna() & tokens.str.contains(r"\d{4}", na=False)
    still = still & ~tokens.str.fullmatch(r"\d{8}", na=False)
    if still.any():
        result.loc[still] = pd.to_datetime(tokens.loc[still], errors="coerce")

    n_empty = int(empty_mask.sum())
    n_invalid = int((~empty_mask & result.isna()).sum())
    return result, n_invalid, n_empty


def voc_period_text(df: pd.DataFrame) -> str:
    dates = df["\uc77c\uc790"].dropna()
    if len(dates) == 0:
        return "-"
    return f"{dates.min().strftime('%Y-%m-%d')} ~ {dates.max().strftime('%Y-%m-%d')}"




# ---------------------------------------------------------------------------
# 유틸
# ---------------------------------------------------------------------------
def now_ts() -> str:
    return datetime.now().strftime("%H:%M:%S")


def stamp() -> str:
    return datetime.now().strftime("%Y%m%d_%H%M%S")


def resolve_file_path(file_obj: Any) -> str | None:
    if file_obj is None:
        return None
    if isinstance(file_obj, str) and file_obj.strip():
        return file_obj
    if isinstance(file_obj, dict):
        return file_obj.get("path") or file_obj.get("name")
    for attr in ("path", "name"):
        value = getattr(file_obj, attr, None)
        if value:
            return str(value)
    return str(file_obj)


def sanitize_log(msg: str) -> str:
    text = str(msg)
    text = re.sub(r"sk-[A-Za-z0-9_-]{8,}", "[REDACTED]", text)
    text = re.sub(r"(?i)(api[_-]?key\s*[:=]\s*)\S+", r"\1[REDACTED]", text)
    text = re.sub(r"(?i)(authorization\s*[:=]\s*)\S+", r"\1[REDACTED]", text)
    text = re.sub(r"(?i)(bearer\s+)\S+", r"\1[REDACTED]", text)
    return text


def set_runtime_api_key(key: str | None) -> None:
    cleaned = (key or "").strip()
    _RUNTIME.api_key = cleaned or None


def resolve_api_key() -> str:
    ui_key = getattr(_RUNTIME, "api_key", None)
    if ui_key:
        return str(ui_key)
    return (os.getenv("OPENAI_API_KEY") or "").strip()


def create_llm(api_key: str) -> LLM:
    cleaned = (api_key or "").strip()
    if not cleaned:
        raise RuntimeError("api-key-missing")
    return LLM(
        model=MODEL_NAME,
        api_key=cleaned,
        temperature=0.2,
        timeout=180,
    )


def esc(value: Any) -> str:
    return html.escape("" if value is None else str(value), quote=True)


def require_df() -> pd.DataFrame:
    with STORE_LOCK:
        df = STORE.get("df")
    if df is None or df.empty:
        raise ValueError("먼저 [파일 업로드]에서 VOC CSV를 불러오세요.")
    return df


def load_voc_csv(path: str | Path) -> pd.DataFrame:
    path = Path(path)
    if not path.exists():
        raise FileNotFoundError("upload-missing")
    if path.suffix.lower() != ".csv":
        raise ValueError("not-csv")
    if path.stat().st_size == 0:
        raise ValueError("empty-file")
    last_error: Exception | None = None
    df: pd.DataFrame | None = None
    for encoding in ("utf-8-sig", "utf-8", "cp949", "euc-kr"):
        try:
            df = pd.read_csv(path, encoding=encoding)
            break
        except UnicodeDecodeError as exc:
            last_error = exc
            continue
        except pd.errors.EmptyDataError as exc:
            raise ValueError("empty-file") from exc
        except pd.errors.ParserError as exc:
            raise ValueError("bad-csv") from exc
    if df is None:
        raise ValueError(f"CSV 인코딩을 읽지 못했습니다: {last_error}")
    if len(df.columns) == 0:
        raise ValueError("empty-file")

    df.columns = [str(c).strip() for c in df.columns]
    missing = [c for c in REQUIRED_COLUMNS if c not in df.columns]
    if missing:
        raise ValueError(
            "필수 컬럼이 없습니다: "
            + ", ".join(missing)
            + f"\n현재 컬럼: {', '.join(map(str, df.columns))}"
        )
    df = df.copy()
    parsed_dates, n_invalid, n_empty = parse_voc_dates(df["일자"])
    df["일자"] = parsed_dates
    df.attrs["date_invalid"] = n_invalid
    df.attrs["date_empty"] = n_empty
    for col in ["고객명", "산업군", "지역", "제품명", "분야", "불만"]:
        df[col] = df[col].fillna("").astype(str).str.strip()
    df = df[df["불만"].str.len() > 0].reset_index(drop=True)
    if df.empty:
        raise ValueError("유효한 불만 내용이 있는 행이 없습니다.")
    return df


def ratio_table(df: pd.DataFrame, column: str) -> pd.DataFrame:
    series = df[column].replace("", "미기재")
    counts = series.value_counts()
    total = int(counts.sum())
    rows = []
    for name, count in counts.items():
        pct = round(count / total * 100, 2)
        rows.append({"항목": str(name), "건수": int(count), "비율(%)": f"{pct:.2f}"})
    return pd.DataFrame(rows)


def save_ratio_png_matplotlib(table: pd.DataFrame, column: str, path: Path) -> Path | None:
    """Chrome/Kaleido 없이 보고서용 PNG를 저장합니다. 실패하면 None을 반환합니다."""
    try:
        names = [str(v) for v in table["항목"].tolist()]
        values = [float(v) for v in table["비율(%)"].tolist()]
        if FONT_PATH:
            fp_title = font_manager.FontProperties(fname=FONT_PATH, size=16)
            fp_axis = font_manager.FontProperties(fname=FONT_PATH, size=11)
        else:
            fp_title = font_manager.FontProperties(family="sans-serif", size=16)
            fp_axis = font_manager.FontProperties(family="sans-serif", size=11)
        fig, ax = plt.subplots(figsize=(12, 6.2))
        bars = ax.bar(names, values, color="#2F6FB3")
        ax.set_title(f"{column}별 VOC 비율", fontproperties=fp_title, color="#1F3A5F")
        ax.set_xlabel(column, fontproperties=fp_axis)
        ax.set_ylabel("비율(%)", fontproperties=fp_axis)
        ax.set_ylim(0, max(values) * 1.18 if values else 1)
        for label in ax.get_xticklabels() + ax.get_yticklabels():
            label.set_fontproperties(fp_axis)
        plt.xticks(rotation=18, ha="right")
        for bar, val in zip(bars, values):
            ax.text(
                bar.get_x() + bar.get_width() / 2,
                bar.get_height(),
                f"{val:.2f}%",
                ha="center",
                va="bottom",
                fontproperties=fp_axis,
            )
        fig.tight_layout()
        fig.savefig(path, dpi=160, facecolor="white", bbox_inches="tight")
        plt.close(fig)
        if path.exists() and path.stat().st_size > 0:
            return register_export_file(path)
    except Exception:
        plt.close("all")
    return None


def make_ratio_bar(df: pd.DataFrame, column: str) -> tuple[pd.DataFrame, go.Figure, Path | None]:
    table = ratio_table(df, column)
    x = table["항목"].tolist()
    y = [float(v) for v in table["비율(%)"].tolist()]
    counts = table["건수"].tolist()
    fig = go.Figure(
        go.Bar(
            x=x,
            y=y,
            name="VOC 비율(%)",
            text=[f"{v:.2f}%" for v in y],
            textposition="outside",
            marker_color="#2F6FB3",
            hovertemplate="%{x}<br>비율: %{y:.2f}%<br>건수: %{customdata}건<extra></extra>",
            customdata=counts,
        )
    )
    fig.update_layout(
        title=dict(text=f"{column}별 VOC 비율", font=dict(family=FONT_FAMILY, size=20, color="#1F3A5F")),
        xaxis=dict(
            title=column,
            tickfont=dict(family=FONT_FAMILY, size=12),
            title_font=dict(family=FONT_FAMILY, size=14),
        ),
        yaxis=dict(
            title="비율(%)",
            tickfont=dict(family=FONT_FAMILY, size=12),
            title_font=dict(family=FONT_FAMILY, size=14),
            rangemode="tozero",
        ),
        font=dict(family=FONT_FAMILY, size=13, color="#222222"),
        template="plotly_white",
        margin=dict(l=50, r=30, t=70, b=80),
        height=520,
        uniformtext_minsize=10,
        uniformtext_mode="show",
        showlegend=True,
        legend=dict(title_text="지표"),
    )
    fig.update_yaxes(range=[0, max(y) * 1.18 if y else 1])
    png_path = EXPORT_DIR / f"voc_ratio_{column}_{stamp()}.png"
    saved = save_ratio_png_matplotlib(table, column, png_path)
    with STORE_LOCK:
        STORE["charts"][column] = saved
        STORE["stats"][column] = table
        STORE["charts"][f"{column}_fig"] = fig
        if saved is None:
            STORE["chart_errors"] = dict(STORE.get("chart_errors") or {})
            STORE["chart_errors"][column] = "비율 그림을 PNG로 저장하지 못했습니다."
    return table, fig, saved


_TOKEN_SPLIT = re.compile(r"[^가-힣A-Za-z0-9+\-]+")
_JOSA_SUFFIXES = (
    "으로는", "으로서", "에서의", "에게서", "으로부터",
    "입니다", "습니다", "됩니다", "합니다", "했습니다", "되었습니다",
    "으로", "에서", "에게", "부터", "까지", "이나", "거나",
    "이", "가", "은", "는", "을", "를", "에", "의", "와", "과", "로", "만", "도",
)


def _strip_josa(token: str) -> str:
    current = token
    changed = True
    while changed and len(current) >= 2:
        changed = False
        for suffix in _JOSA_SUFFIXES:
            if current.endswith(suffix) and len(current) - len(suffix) >= 2:
                current = current[: -len(suffix)]
                changed = True
                break
    return current


def extract_keywords(texts: list[str], top_n: int = 80) -> dict[str, int]:
    counter: Counter[str] = Counter()
    for text in texts:
        for raw in _TOKEN_SPLIT.split(str(text)):
            if not raw:
                continue
            if re.match(r"^[A-Za-z]", raw):
                key = raw.upper() if len(raw) <= 5 else raw
            else:
                key = _strip_josa(raw)
            if len(key) < 2 or key in KOREAN_STOPWORDS:
                continue
            if not re.fullmatch(r"[가-힣]+|[A-Za-z][A-Za-z0-9+\-/]*", key):
                continue
            counter[key] += 1
    if not counter:
        return {"데이터없음": 1}
    return dict(counter.most_common(top_n))


def make_wordcloud(df: pd.DataFrame) -> tuple[Path, dict[str, int]]:
    if not FONT_PATH:
        raise FileNotFoundError(FONT_WARNING or "missing-korean-font")
    freq = extract_keywords(df["불만"].tolist())
    wc = WordCloud(
        font_path=FONT_PATH,
        width=1400,
        height=800,
        background_color="white",
        max_words=120,
        collocations=False,
        prefer_horizontal=0.92,
        min_font_size=10,
        colormap="tab20",
        regexp=r"[가-힣A-Za-z0-9+\-]+",
    ).generate_from_frequencies(freq)
    path = EXPORT_DIR / f"voc_wordcloud_{stamp()}.png"
    plt.figure(figsize=(14, 8))
    plt.imshow(wc, interpolation="bilinear")
    plt.axis("off")
    plt.tight_layout(pad=0)
    plt.savefig(path, dpi=160, bbox_inches="tight", facecolor="white")
    plt.close()
    register_export_file(path)
    with STORE_LOCK:
        STORE["wordcloud_path"] = path
        STORE["keywords"] = freq
    return path, freq


def voc_overview(df: pd.DataFrame) -> str:
    period = voc_period_text(df)
    n_invalid = int(df.attrs.get("date_invalid", 0) or 0)
    extra = f" / 일자 해석 실패 {n_invalid}건은 기간에서 제외" if n_invalid else ""
    return (
        f"총 VOC {len(df):,}건 / 기간 {period}{extra} / "
        f"산업군 {df['산업군'].nunique()}개 / 제품 {df['제품명'].nunique()}개 / "
        f"분야 {df['분야'].nunique()}개"
    )


def build_voc_digest(df: pd.DataFrame, max_industries: int = 8, per_industry: int = 4) -> str:
    chunks: list[str] = [voc_overview(df), ""]
    for col in ["산업군", "분야", "제품명"]:
        table = ratio_table(df, col)
        chunks.append(f"[{col}별 통계]")
        for _, row in table.iterrows():
            chunks.append(f"- {row['항목']}: {row['건수']}건 ({row['비율(%)']}%)")
        chunks.append("")
    chunks.append("[산업군별 대표 불만]")
    for industry, count in df["산업군"].value_counts().head(max_industries).items():
        chunks.append(f"## {industry} ({int(count)}건)")
        samples = df.loc[df["산업군"] == industry, "불만"].head(per_industry).tolist()
        for sample in samples:
            chunks.append(f"- {sample}")
        fields = ratio_table(df[df["산업군"] == industry], "분야")
        top_fields = ", ".join(
            f"{r['항목']} {r['비율(%)']}%" for _, r in fields.head(4).iterrows()
        )
        chunks.append(f"  분야 비중: {top_fields}")
        chunks.append("")
    keywords = extract_keywords(df["불만"].tolist(), top_n=25)
    chunks.append("[불만 키워드 Top]")
    chunks.append(", ".join(f"{k}({v})" for k, v in keywords.items()))
    return "\n".join(chunks)


# ---------------------------------------------------------------------------
# 워드 보고서 (한글 폰트)
# ---------------------------------------------------------------------------
def _set_run_font(run, size_pt: float = 11, bold: bool = False, color: RGBColor | None = None) -> None:
    run.bold = bold
    run.font.size = Pt(size_pt)
    run.font.color.rgb = color or DARK
    run.font.name = FONT_FAMILY
    rPr = run._element.get_or_add_rPr()
    rFonts = rPr.find(qn("w:rFonts"))
    if rFonts is None:
        rFonts = OxmlElement("w:rFonts")
        rPr.append(rFonts)
    rFonts.set(qn("w:ascii"), FONT_FAMILY)
    rFonts.set(qn("w:hAnsi"), FONT_FAMILY)
    rFonts.set(qn("w:eastAsia"), FONT_EAST_ASIA)
    lang = rPr.find(qn("w:lang"))
    if lang is None:
        lang = OxmlElement("w:lang")
        rPr.append(lang)
    lang.set(qn("w:val"), "ko-KR")
    lang.set(qn("w:eastAsia"), "ko-KR")


def _style_font(style, size_pt: float = 11) -> None:
    style.font.name = FONT_FAMILY
    style.font.size = Pt(size_pt)
    rPr = style.element.get_or_add_rPr()
    rFonts = rPr.find(qn("w:rFonts"))
    if rFonts is None:
        rFonts = OxmlElement("w:rFonts")
        rPr.append(rFonts)
    rFonts.set(qn("w:ascii"), FONT_FAMILY)
    rFonts.set(qn("w:hAnsi"), FONT_FAMILY)
    rFonts.set(qn("w:eastAsia"), FONT_EAST_ASIA)


def _shade_cell(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    shd.set(qn("w:val"), "clear")
    tc_pr.append(shd)


def add_kr_paragraph(
    doc: Document,
    text: str,
    size: float = 11,
    bold: bool = False,
    color: RGBColor | None = None,
    align: str = "left",
    space_after: float = 6,
) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
    if align == "center":
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    elif align == "right":
        p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = p.add_run(text)
    _set_run_font(run, size_pt=size, bold=bold, color=color)


def add_kr_heading(doc: Document, text: str, level: int = 1) -> None:
    sizes = {0: 20, 1: 16, 2: 13, 3: 12}
    colors = {0: NAVY, 1: NAVY, 2: ACCENT, 3: DARK}
    add_kr_paragraph(
        doc,
        text,
        size=sizes.get(level, 12),
        bold=True,
        color=colors.get(level, DARK),
        space_after=8 if level <= 1 else 4,
    )


def add_stats_table(doc: Document, table_df: pd.DataFrame) -> None:
    rows, cols = table_df.shape
    table = doc.add_table(rows=rows + 1, cols=cols)
    table.style = "Table Grid"
    headers = list(table_df.columns)
    for j, header in enumerate(headers):
        cell = table.rows[0].cells[j]
        cell.text = ""
        p = cell.paragraphs[0]
        run = p.add_run(str(header))
        _set_run_font(run, size_pt=10, bold=True, color=RGBColor(255, 255, 255))
        _shade_cell(cell, "1F3A5F")
    for i, row in enumerate(table_df.itertuples(index=False), start=1):
        for j, value in enumerate(row):
            cell = table.rows[i].cells[j]
            cell.text = ""
            p = cell.paragraphs[0]
            run = p.add_run(str(value))
            _set_run_font(run, size_pt=10)
            if i % 2 == 0:
                _shade_cell(cell, "F4F7FB")
    doc.add_paragraph()


def add_picture_if_exists(doc: Document, path: Path | None) -> bool:
    if path and Path(path).exists():
        try:
            doc.add_picture(str(path), width=Cm(16.0))
            last = doc.paragraphs[-1]
            last.alignment = WD_ALIGN_PARAGRAPH.CENTER
            last.paragraph_format.space_after = Pt(10)
            return True
        except Exception:
            return False
    return False


def write_markdown_like(doc: Document, text: str) -> None:
    if not text or not str(text).strip():
        add_kr_paragraph(doc, "분석 결과가 없습니다. 원본 VOC를 다시 확인하세요.", color=RGBColor(0x88, 0x00, 0x00))
        return
    for raw in str(text).replace("\r\n", "\n").split("\n"):
        line = raw.rstrip()
        if not line.strip():
            continue
        stripped = line.strip()
        if stripped.startswith("### "):
            add_kr_heading(doc, stripped[4:].strip(), level=3)
        elif stripped.startswith("## "):
            add_kr_heading(doc, stripped[3:].strip(), level=2)
        elif stripped.startswith("# "):
            add_kr_heading(doc, stripped[2:].strip(), level=1)
        elif re.match(r"^[-*•]\s+", stripped):
            add_kr_paragraph(doc, "· " + re.sub(r"^[-*•]\s+", "", stripped), size=11, space_after=3)
        elif re.match(r"^\d+[.)]\s+", stripped):
            add_kr_paragraph(doc, stripped, size=11, space_after=3)
        else:
            add_kr_paragraph(doc, stripped, size=11, space_after=4)


def generate_word_report(issue_analysis: str, missing_images: list[str] | None = None) -> Path:
    df = require_df()
    with STORE_LOCK:
        industry_table = STORE["stats"].get("산업군")
        field_table = STORE["stats"].get("분야")
        industry_chart = STORE["charts"].get("산업군")
        field_chart = STORE["charts"].get("분야")
        wc_path = STORE.get("wordcloud_path")

    if industry_table is None:
        industry_table = ratio_table(df, "산업군")
    if field_table is None:
        field_table = ratio_table(df, "분야")
    if wc_path is not None and not Path(wc_path).exists():
        wc_path = None

    doc = Document()
    section = doc.sections[0]
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.left_margin = Cm(2.3)
    section.right_margin = Cm(2.3)
    section.top_margin = Cm(2.2)
    section.bottom_margin = Cm(2.2)

    for style_name in ["Normal", "Title", "Heading 1", "Heading 2", "Heading 3"]:
        try:
            _style_font(doc.styles[style_name], 11)
        except KeyError:
            pass

    omitted_early = list(missing_images or [])
    title = "고객사 VOC 분석 보고서"
    if omitted_early:
        title = "고객사 VOC 분석 보고서 (그림 누락 · 불완전)"
    add_kr_heading(doc, title, level=0)
    add_kr_paragraph(doc, voc_overview(df), size=11, color=ACCENT, space_after=2)
    add_kr_paragraph(
        doc,
        f"작성일 {datetime.now().strftime('%Y-%m-%d %H:%M')}  |  통계·그림은 이 컴퓨터에서 계산",
        size=10,
        color=RGBColor(0x66, 0x66, 0x66),
        space_after=12,
    )
    add_kr_paragraph(
        doc,
        "이 보고서는 업로드된 VOC를 바탕으로 산업군·분야 분포와 불만 키워드를 정리하고, "
        "의사결정에 필요한 주요 이슈와 대응 방향을 제시합니다.",
        size=11,
        space_after=14,
    )

    omitted: list[str] = list(missing_images or [])

    add_kr_heading(doc, "1. 산업군별 VOC 통계", level=1)
    add_kr_paragraph(doc, "산업군 기준 VOC 건수와 비중(소수점 둘째 자리)입니다.", size=11)
    add_stats_table(doc, industry_table)
    if not add_picture_if_exists(doc, industry_chart):
        omitted.append("산업군 비율 그림")
        add_kr_paragraph(doc, "산업군 비율 그림을 이번 보고서에 넣지 못했습니다.", color=RGBColor(0x88, 0x00, 0x00))

    add_kr_heading(doc, "2. 분야별 VOC 통계", level=1)
    add_kr_paragraph(doc, "품질·납기·기술지원 등 분야 기준 VOC 비중입니다.", size=11)
    add_stats_table(doc, field_table)
    if not add_picture_if_exists(doc, field_chart):
        omitted.append("분야 비율 그림")
        add_kr_paragraph(doc, "분야 비율 그림을 이번 보고서에 넣지 못했습니다.", color=RGBColor(0x88, 0x00, 0x00))

    add_kr_heading(doc, "3. 불만 키워드 워드클라우드", level=1)
    add_kr_paragraph(doc, "불만 텍스트에서 추출한 핵심 키워드의 상대 빈도입니다.", size=11)
    if not add_picture_if_exists(doc, wc_path):
        omitted.append("워드클라우드")
        add_kr_paragraph(doc, "워드클라우드 그림을 이번 보고서에 넣지 못했습니다.", color=RGBColor(0x88, 0x00, 0x00))

    add_kr_heading(doc, "4. 주요 이슈", level=1)
    issue_body = issue_analysis or ""
    # 대응방안 섹션은 5장으로 분리
    split_match = re.split(r"(?m)^(?:#+\s*)?(개선\s*과제|대응\s*방안|권고\s*사항)", issue_body, maxsplit=1)
    if len(split_match) >= 3:
        main_issue = split_match[0].strip()
        rest_title = split_match[1].strip()
        rest_body = split_match[2].strip()
        write_markdown_like(doc, main_issue)
        add_kr_heading(doc, f"5. {rest_title}", level=1)
        write_markdown_like(doc, rest_body)
    else:
        write_markdown_like(doc, issue_body)
        add_kr_heading(doc, "5. 개선과제 및 대응방안", level=1)
        add_kr_paragraph(doc, "이슈 분석 결과에 대응방안이 포함되지 않았습니다. Issue Agent 출력을 확인하세요.")

    unique_omitted = list(dict.fromkeys(omitted))
    if unique_omitted:
        add_kr_paragraph(
            doc,
            "이 파일은 그림이 빠져 있어 완전한 보고서가 아닙니다. 넣지 못한 항목: "
            + ", ".join(unique_omitted),
            size=10,
            color=RGBColor(0x88, 0x00, 0x00),
            space_after=8,
        )

    add_kr_paragraph(
        doc,
        "※ 수치는 업로드 CSV에서 계산한 값이며, 자료에 없는 건수·일정·담당자는 단정하지 않았습니다.",
        size=9,
        color=RGBColor(0x66, 0x66, 0x66),
        space_after=0,
    )

    out_path = EXPORT_DIR / f"VOC분석보고서_{stamp()}.docx"
    doc.save(str(out_path))
    register_export_file(out_path)
    with STORE_LOCK:
        STORE["report_path"] = out_path
        STORE["issue_analysis"] = issue_analysis
        STORE["report_missing_images"] = unique_omitted
    return out_path


# ---------------------------------------------------------------------------
# CrewAI 도구
# ---------------------------------------------------------------------------
@tool("analyze_voc_ratio")
def analyze_voc_ratio_tool(column: str) -> str:
    """선택한 컬럼(산업군/제품명/분야)별 VOC 건수와 비율(%)을 계산하고 막대그래프 PNG를 저장합니다."""
    df = require_df()
    column = str(column).strip()
    if column not in STAT_COLUMNS:
        return f"column은 {STAT_COLUMNS} 중 하나여야 합니다. 입력값: {column}"
    table, _, png_path = make_ratio_bar(df, column)
    lines = [f"{column}별 VOC 비율 그래프 저장: {png_path}"]
    for _, row in table.iterrows():
        lines.append(f"- {row['항목']}: {row['건수']}건, {row['비율(%)']}%")
    return "\n".join(lines)


@tool("create_voc_wordcloud")
def create_voc_wordcloud_tool() -> str:
    """불만 컬럼에서 한글 키워드를 추출하고 워드클라우드 PNG를 저장합니다."""
    df = require_df()
    path, freq = make_wordcloud(df)
    top = ", ".join(f"{k}({v})" for k, v in list(freq.items())[:20])
    return f"워드클라우드 저장: {path}\n상위 키워드: {top}"


@tool("create_voc_word_report")
def create_voc_word_report_tool(issue_analysis: str) -> str:
    """산업군/분야 통계, 워드클라우드, 주요이슈, 대응방안을 담은 한글 워드 보고서를 생성합니다."""
    path = generate_word_report(issue_analysis)
    return f"워드 보고서 저장: {path}"


# ---------------------------------------------------------------------------
# Agent / Crew
# ---------------------------------------------------------------------------
def make_voc_agent(llm: LLM) -> Agent:
    return Agent(
        role="VOC 데이터 분석가",
        goal="고객 VOC CSV를 정확히 읽고 항목별 비율 통계, 시각화, 워드클라우드를 생성한다.",
        backstory=(
            "제조·금융·유통 등 고객사 VOC를 다루는 데이터 분석가입니다. "
            "숫자는 원본 CSV에서만 계산하고, 비율은 소수점 둘째 자리까지 표기합니다. "
            "시각화와 워드클라우드는 제공된 도구만 사용합니다."
        ),
        llm=llm,
        tools=[analyze_voc_ratio_tool, create_voc_wordcloud_tool],
        verbose=True,
        allow_delegation=False,
        max_iter=6,
        cache=False,
    )


def make_issue_agent(llm: LLM) -> Agent:
    return Agent(
        role="VOC 이슈 도출 전문가",
        goal="산업군별 주요 VOC를 분석해 핵심 이슈와 개선과제를 도출한다.",
        backstory=(
            "품질·CS·사업 담당 임원이 5분 안에 판단할 수 있도록 쓰는 이슈 분석가입니다. "
            "사실(Fact)만 데이터에 두고, 이슈(Issue)와 권고(Recommend)를 연결합니다. "
            "자료에 없는 수치·일정·담당자를 만들지 않습니다."
        ),
        llm=llm,
        tools=[],
        verbose=True,
        allow_delegation=False,
        max_iter=8,
        cache=False,
    )


def make_report_agent(llm: LLM) -> Agent:
    return Agent(
        role="VOC 보고서 작성가",
        goal="VOC 통계와 이슈 분석 결과를 한글 워드 보고서로 만든다.",
        backstory=(
            "임원 보고용 한글 문서를 작성하는 보고 담당자입니다. "
            "통계 표·그래프·워드클라우드·주요이슈·대응방안이 한 파일에 들어가도록 "
            "보고서 생성 도구를 반드시 호출합니다."
        ),
        llm=llm,
        tools=[create_voc_word_report_tool],
        verbose=True,
        allow_delegation=False,
        max_iter=6,
        cache=False,
    )


class _FnLogHandler(logging.Handler):
    def __init__(self, fn: Callable[[str], None]):
        super().__init__()
        self.fn = fn

    def emit(self, record: logging.LogRecord) -> None:
        try:
            msg = self.format(record).strip()
            if msg:
                self.fn(f"[CrewAI] {sanitize_log(msg)[:400]}")
        except Exception:
            return


def run_crew(agents: list[Agent], tasks: list[Task], logger: Callable[[str], None], name: str) -> Any:
    logger(f"{name} Crew 실행 (모델: {MODEL_NAME})")
    handler = _FnLogHandler(logger)
    handler.setFormatter(logging.Formatter("%(message)s"))
    crew_logger = logging.getLogger("crewai")
    prev_level = crew_logger.level
    crew_logger.addHandler(handler)
    crew_logger.setLevel(logging.INFO)

    def _step(output: Any) -> None:
        text = str(output)
        logger("진행 중: " + text.replace("\n", " ")[:220])

    crew = Crew(
        agents=agents,
        tasks=tasks,
        process=Process.sequential,
        verbose=True,
        memory=False,
        tracing=False,
        cache=False,
        step_callback=_step,
    )
    buf = io.StringIO()
    try:
        with contextlib_redirect(buf):
            result = crew.kickoff()
    finally:
        crew_logger.removeHandler(handler)
        crew_logger.setLevel(prev_level)

    for line in buf.getvalue().splitlines():
        line = line.strip()
        if line:
            logger(line[:300])
    logger(f"{name} Crew 완료")
    return result


class contextlib_redirect:
    """stdout/stderr를 버퍼로 옮깁니다. rich Console 출력도 가능한 범위에서 수집합니다."""

    def __init__(self, buf: io.StringIO):
        self.buf = buf
        self._stdout = None
        self._stderr = None

    def __enter__(self):
        self._stdout, self._stderr = sys.stdout, sys.stderr
        sys.stdout = _Tee(self._stdout, self.buf)
        sys.stderr = _Tee(self._stderr, self.buf)
        return self.buf

    def __exit__(self, *args):
        sys.stdout = self._stdout
        sys.stderr = self._stderr


class _Tee(io.TextIOBase):
    def __init__(self, original, buf: io.StringIO):
        self.original = original
        self.buf = buf

    def write(self, s):
        try:
            self.original.write(s)
        except Exception:
            pass
        try:
            self.buf.write(s)
        except Exception:
            pass
        return len(s) if isinstance(s, str) else 0

    def flush(self):
        try:
            self.original.flush()
        except Exception:
            pass


def raw_text(result: Any) -> str:
    if result is None:
        return ""
    for attr in ("raw", "output"):
        value = getattr(result, attr, None)
        if value:
            return str(value)
    return str(result)


# ---------------------------------------------------------------------------
# 메뉴별 작업
# ---------------------------------------------------------------------------

JOB_LOCK = threading.Lock()
JobFn = Callable[[Callable[[str], None]], Any]


def persist_user_csv(file_obj: Any) -> tuple[Path, str]:
    src = resolve_file_path(file_obj)
    if not src:
        raise FileNotFoundError("upload-missing")
    src_path = Path(src)
    if not src_path.exists():
        raise FileNotFoundError("upload-missing")
    if src_path.suffix.lower() != ".csv":
        raise ValueError("not-csv")
    if src_path.stat().st_size == 0:
        raise ValueError("empty-file")
    raw = src_path.read_bytes()
    name = src_path.name
    dest = UPLOAD_DIR / f"upload_{stamp()}_{name}"
    dest.write_bytes(raw)
    return dest, name


def friendly_error(exc: BaseException) -> str:
    text = sanitize_log(str(exc))
    low = text.lower()
    if isinstance(exc, FileNotFoundError) or "upload-missing" in text:
        return "선택한 CSV를 찾지 못했습니다. 파일을 다시 선택한 뒤 적용해 주세요."
    if "not-csv" in text:
        return "지원 형식은 CSV입니다. .csv 파일을 선택해 주세요."
    if "empty-file" in text:
        return "파일이 비어 있습니다. 열 이름이 있는 VOC CSV를 선택해 주세요."
    if "bad-csv" in text:
        return "CSV 형식이 올바르지 않습니다. 쉼표로 구분된 표를 확인해 주세요."
    if "필수 컬럼" in text:
        return "필수 열이 없는 파일입니다. 순번, 일자, 고객명, 산업군, 지역, 제품명, 분야, 불만 열을 확인해 주세요."
    if "인코딩" in text:
        return "파일 글자를 읽지 못했습니다. UTF-8 CSV로 저장한 뒤 다시 시도해 주세요."
    if "유효한 불만" in text:
        return "불만 내용이 있는 행이 없습니다. 다른 CSV를 선택해 주세요."
    if "missing-korean-font" in text or "한글 폰트" in text:
        return (
            "한글 폰트를 찾지 못했습니다. Windows 맑은 고딕이나 VOC_FONT_PATH에 TTF 경로를 지정한 뒤 다시 시도해 주세요."
        )
    if "먼저" in text and "업로드" in text:
        return "아직 적용된 VOC 데이터가 없습니다. 데이터 불러오기에서 CSV를 적용하거나 실습 샘플을 선택해 주세요."
    if "통계 항목" in text:
        return "통계 기준은 산업군, 제품명, 분야 중에서 고를 수 있습니다."
    if "api-key-missing" in text or "openai_api_key" in low:
        return (
            "OpenAI API 키가 없습니다. 화면에서 키를 입력하거나 환경 변수 OPENAI_API_KEY를 설정하세요. "
            "통계 막대그래프와 워드클라우드는 키 없이 이 컴퓨터에서 만들 수 있습니다."
        )
    if (
        "incorrect api key" in low
        or "invalid_api_key" in low
        or "invalid api key" in low
        or "unauthorized" in low
        or " 401" in f" {text}"
        or "status code: 401" in low
    ):
        return "OpenAI 인증에 실패했습니다. 키를 확인한 뒤 다시 시도해 주세요."
    if "429" in text or "rate limit" in low or "insufficient_quota" in low:
        return "OpenAI 요청 한도 또는 할당량을 넘었습니다. 잠시 후 다시 시도해 주세요."
    if "timeout" in low or "timed out" in low:
        return "OpenAI 요청 시간이 초과되었습니다. 잠시 후 다시 시도해 주세요."
    if "connection" in low or "connecterror" in low or "network" in low:
        return "OpenAI 서버에 연결하지 못했습니다. 인터넷 연결을 확인한 뒤 다시 시도해 주세요."
    return "요청을 처리하지 못했습니다. 파일을 확인한 뒤 다시 시도해 주세요."


def empty_state() -> dict[str, Any]:
    return {
        "source_kind": None,
        "source_label": "데이터 없음",
        "row_count": 0,
        "overview": "",
        "period": "-",
        "n_industry": 0,
        "n_product": 0,
        "n_field": 0,
        "status": "idle",
        "status_text": "대기",
        "view_title": "데이터 불러오기",
        "busy": False,
        "has_data": False,
        "last_error": "",
        "stats_column": None,
        "stats_ready": False,
        "wc_ready": False,
        "report_ready": False,
    }


def preview_frame(df: pd.DataFrame) -> pd.DataFrame:
    out = df.copy()
    if "일자" in out.columns and pd.api.types.is_datetime64_any_dtype(out["일자"]):
        out["일자"] = out["일자"].dt.strftime("%Y-%m-%d").fillna("")
    complaints = out["불만"].astype(str)
    out["불만 미리보기"] = [s if len(s) <= 48 else s[:48] + "…" for s in complaints.tolist()]
    cols = [c for c in ["순번", "일자", "고객명", "산업군", "지역", "제품명", "분야", "불만 미리보기"] if c in out.columns]
    out = out[cols].copy()
    return out.astype(str)


def full_frame(df: pd.DataFrame) -> pd.DataFrame:
    out = df.copy()
    if "일자" in out.columns and pd.api.types.is_datetime64_any_dtype(out["일자"]):
        out["일자"] = out["일자"].dt.strftime("%Y-%m-%d").fillna("")
    return out.astype(str)


EMPTY_PREVIEW = pd.DataFrame(
    columns=["순번", "일자", "고객명", "산업군", "지역", "제품명", "분야", "불만 미리보기"]
)
EMPTY_FULL = pd.DataFrame(columns=["순번", "일자", "고객명", "산업군", "지역", "제품명", "분야", "불만"])
EMPTY_STATS = pd.DataFrame(columns=["구분", "VOC 건수", "비율(%)"])


def display_stats_table(table: pd.DataFrame) -> pd.DataFrame:
    return table.rename(columns={"항목": "구분", "건수": "VOC 건수"})


def snapshot_from_df(df: pd.DataFrame, source_kind: str, source_label: str) -> dict[str, Any]:
    period = voc_period_text(df)
    state = empty_state()
    state.update(
        {
            "source_kind": source_kind,
            "source_label": source_label,
            "row_count": int(len(df)),
            "overview": voc_overview(df),
            "period": period,
            "n_industry": int(df["산업군"].nunique()),
            "n_product": int(df["제품명"].nunique()),
            "n_field": int(df["분야"].nunique()),
            "status": "ready",
            "status_text": "적용됨",
            "has_data": True,
            "busy": False,
            "view_title": "데이터 불러오기",
        }
    )
    return state


def status_badge(state: dict[str, Any]) -> str:
    kind = state.get("status") or "idle"
    mapping = {
        "idle": "badge-idle",
        "uploading": "badge-busy",
        "analyzing": "badge-busy",
        "ready": "badge-ready",
        "done": "badge-ready",
        "failed": "badge-fail",
    }
    cls = mapping.get(kind, "badge-idle")
    label = esc(state.get("status_text") or "대기")
    return f'<span class="badge {cls}">{label}</span>'


def render_header(state: dict[str, Any]) -> str:
    label = esc(state.get("source_label") or "데이터 없음")
    if state.get("source_kind") == "sample":
        kind = "실습 샘플"
    elif state.get("source_kind") == "file":
        kind = "업로드 파일"
    else:
        kind = "없음"
    count_txt = f"{state.get('row_count') or 0:,}건" if state.get("has_data") else "-"
    return f"""
    <div class="app-header">
      <div>
        <div class="app-kicker">고객사 VOC 분석</div>
        <div class="app-title">{esc(state.get("view_title") or "데이터 불러오기")}</div>
      </div>
      <div class="app-header-meta">
        <div class="meta-block">
          <div class="meta-k">현재 데이터</div>
          <div class="meta-v">{label}</div>
          <div class="meta-s">{kind}</div>
        </div>
        <div class="meta-block">
          <div class="meta-k">건수</div>
          <div class="meta-v">{esc(count_txt)}</div>
          <div class="meta-s">{esc(state.get("period") or "-")}</div>
        </div>
        <div class="meta-block">
          <div class="meta-k">진행 상태</div>
          <div class="meta-v">{status_badge(state)}</div>
          <div class="meta-s">적용된 데이터와 보고 있는 화면</div>
        </div>
      </div>
    </div>
    """


def render_upload_cards(state: dict[str, Any]) -> str:
    if not state.get("has_data"):
        return ""
    return f"""
    <div class="result-head">
      <h2>현재 적용된 VOC</h2>
      <p class="result-sub">{esc(state.get("source_label"))} · {esc(state.get("overview"))}</p>
    </div>
    <div class="cards">
      <div class="card"><div class="k">VOC 건수</div><div class="v">{state["row_count"]:,}</div><div class="s">불만 내용이 있는 행</div></div>
      <div class="card"><div class="k">기간</div><div class="v">{esc(state.get("period"))}</div><div class="s">일자 기준</div></div>
      <div class="card"><div class="k">산업군</div><div class="v">{state.get("n_industry")}</div><div class="s">고유 값 수</div></div>
      <div class="card"><div class="k">제품 / 분야</div><div class="v">{state.get("n_product")} / {state.get("n_field")}</div><div class="s">고유 값 수</div></div>
    </div>
    """


def render_stats_cards(table: pd.DataFrame, column: str, state: dict[str, Any]) -> str:
    if table is None or table.empty:
        return f"""
        <div class="result-head">
          <h2>항목별 VOC 비율</h2>
          <p class="result-sub">분석 대상: {esc(state.get("source_label"))} · 기준: {esc(column)}</p>
        </div>
        """
    top = table.iloc[0]
    n_cat = len(table)
    total = int(table["건수"].sum()) if "건수" in table.columns else state.get("row_count") or 0
    return f"""
    <div class="result-head">
      <h2>항목별 VOC 비율</h2>
      <p class="result-sub">분석 대상: {esc(state.get("source_label"))} · 기준: {esc(column)} · 비율은 소수점 둘째 자리</p>
    </div>
    <div class="cards">
      <div class="card"><div class="k">가장 비중이 큰 구분</div><div class="v">{esc(top["항목"])}</div><div class="s">{esc(top["비율(%)"])}% · {esc(top["건수"])}건</div></div>
      <div class="card"><div class="k">집계 건수</div><div class="v">{total:,}</div><div class="s">현재 적용된 VOC</div></div>
      <div class="card"><div class="k">구분 수</div><div class="v">{n_cat}</div><div class="s">건수가 0인 구분은 목록에 없습니다</div></div>
    </div>
    """


def render_wc_cards(state: dict[str, Any]) -> str:
    freq = STORE.get("keywords") or {}
    top = list(freq.items())[:5]
    chips = "".join(f'<span class="kw">{esc(k)} <em>{esc(v)}</em></span>' for k, v in top)
    if not chips:
        chips = "<span class='muted'>표시할 키워드가 없습니다.</span>"
    return f"""
    <div class="result-head">
      <h2>불만 키워드</h2>
      <p class="result-sub">분석 대상: {esc(state.get("source_label"))} · 불만 열에서 추출한 상대 빈도</p>
    </div>
    <div class="card">
      <div class="k">자주 등장한 키워드</div>
      <div class="kw-row">{chips}</div>
      <div class="s">글자 크기는 빈도를 나타내며, 원인이나 해결책이 확정되었음을 뜻하지 않습니다.</div>
    </div>
    """


def render_notice(kind: str, title: str, body: str) -> str:
    safe_kind = "info" if kind not in {"info", "busy", "fail"} else kind
    return (
        f'<div class="notice notice-{safe_kind}">'
        f'<div class="notice-t">{esc(title)}</div>'
        f'<div class="notice-b">{esc(body)}</div></div>'
    )


INFO_UPLOAD = render_notice(
    "info",
    "데이터를 먼저 적용하세요",
    "CSV를 선택한 뒤 [이 파일 적용]을 누르거나 실습 샘플을 선택하세요. 새 데이터가 적용되면 이전 분석 결과는 비웁니다.",
)
INFO_STATS = render_notice("info", "통계를 아직 계산하지 않았습니다", "데이터가 있으면 기준을 고른 뒤 [비율 계산]을 실행하세요.")
INFO_WC = render_notice("info", "워드클라우드를 아직 만들지 않았습니다", "데이터가 있으면 [키워드 그림 만들기]를 실행하세요.")
INFO_REPORT = render_notice(
    "info",
    "보고서를 아직 만들지 않았습니다",
    "데이터가 있으면 [보고서 만들기]를 실행하세요. 표와 그림·DOCX는 이 컴퓨터에서 만들고, 이슈 문장만 OpenAI API가 있을 때 생성합니다. 대표 불만 문구에 이름 등 개인정보가 있으면 그대로 전송될 수 있습니다.",
)



def current_upload_views(state: dict[str, Any]) -> tuple[pd.DataFrame, pd.DataFrame, str]:
    df = STORE.get("df")
    if df is None or not state.get("has_data"):
        return EMPTY_PREVIEW, EMPTY_FULL, ""
    return preview_frame(df), full_frame(df), render_upload_cards(state)

def sample_button_update(state: dict[str, Any], interactive: bool = True):
    selected = state.get("source_kind") == "sample" and state.get("has_data")
    if selected:
        return gr.update(
            value="선택됨  ·  실습 샘플",
            elem_classes=["source-chip", "is-selected"],
            interactive=interactive,
        )
    return gr.update(
        value="실습 샘플 불러오기",
        elem_classes=["source-chip"],
        interactive=interactive,
    )


def control_updates(busy: bool, state: dict[str, Any], clear_file: bool = False):
    enable = not busy
    file_update = gr.update(interactive=enable)
    if clear_file:
        file_update = gr.update(value=None, interactive=enable)
    return (
        file_update,
        gr.update(interactive=enable),
        sample_button_update(state, interactive=enable),
        gr.update(interactive=enable),
        gr.update(interactive=enable),
        gr.update(interactive=enable),
    )


def apply_loaded_df(df: pd.DataFrame, source_kind: str, source_label: str, logger) -> dict[str, Any]:
    with STORE_LOCK:
        STORE["df"] = df
        STORE["csv_path"] = str(UPLOAD_DIR / "uploaded_voc.csv")
        STORE["chart_errors"] = {}
        STORE["source_kind"] = source_kind
        STORE["source_label"] = source_label
        STORE["charts"] = {}
        STORE["stats"] = {}
        STORE["wordcloud_path"] = None
        STORE["keywords"] = {}
        STORE["issue_analysis"] = ""
        STORE["report_path"] = None
        STORE["report_missing_images"] = []
    n_invalid = int(df.attrs.get("date_invalid", 0) or 0)
    if n_invalid:
        logger(
            f"일자 {n_invalid}건은 지원 형식이 아니어 기간 계산에서 제외했습니다."
        )
    logger(voc_overview(df))
    return snapshot_from_df(df, source_kind, source_label)


def do_upload(path: str | Path, source_kind: str, source_label: str, logger: Callable[[str], None]) -> dict[str, Any]:
    logger("CSV를 읽고 열을 확인합니다.")
    df = load_voc_csv(path)
    save_path = UPLOAD_DIR / "uploaded_voc.csv"
    df.to_csv(save_path, index=False, encoding="utf-8-sig")
    state = apply_loaded_df(df, source_kind, source_label, logger)
    logger("이전 통계·워드클라우드·보고서 결과는 새 데이터와 섞이지 않도록 비웠습니다.")
    return {
        "state": state,
        "preview": preview_frame(df),
        "full": full_frame(df),
        "cards": render_upload_cards(state),
    }


def do_stats(column: str, logger: Callable[[str], None]) -> tuple[Any, pd.DataFrame]:
    df = require_df()
    if column not in STAT_COLUMNS:
        raise ValueError(f"통계 항목은 {STAT_COLUMNS} 중 하나여야 합니다.")
    logger(f"'{column}' 비율을 이 컴퓨터에서 계산합니다. OpenAI는 호출하지 않습니다.")
    table, fig, png = make_ratio_bar(df, column)
    if png:
        logger(f"보고서용 그림 저장: {png.name}")
    else:
        logger("보고서용 PNG는 만들지 못했지만, 화면 그래프와 표는 표시합니다.")
    logger(f"통계 완료 ({len(table)}개 항목)")
    return (fig, table)


def do_wordcloud(logger: Callable[[str], None]) -> tuple[str]:
    df = require_df()
    logger("불만 키워드와 워드클라우드를 이 컴퓨터에서 만듭니다. OpenAI는 호출하지 않습니다.")
    path, freq = make_wordcloud(df)
    logger("상위 키워드: " + ", ".join(list(freq)[:15]))
    logger(f"워드클라우드 완료: {Path(path).name}")
    return (str(path),)


def run_issue_analysis(digest: str, api_key: str, logger: Callable[[str], None]) -> str:
    llm = None
    crew = None
    issue_agent = None
    issue_task = None
    try:
        llm = create_llm(api_key)
        issue_agent = make_issue_agent(llm)
        issue_task = Task(
            description=(
                "당신은 품질/CS 담당 임원이 대응 우선순위를 정하도록 돕는 분석가입니다.\n"
                "VOC 통계 요약과 대표 불만을 보고 산업군별 주요 VOC를 분석한 뒤 "
                "주요 이슈와 개선과제를 도출하세요.\n"
                "규칙: 자료에 없는 수치·고객명·일정을 만들지 말 것. "
                "이슈는 Fact-Issue-Recommend로 연결할 것. "
                "대응방안은 결론-근거-방법-효과 순으로 쓸 것.\n\n"
                "출력 형식(마크다운):\n"
                "# 핵심 메시지\n"
                "# 산업군별 주요 이슈\n"
                "## (산업군명)\n"
                "- 이슈:\n- 근거 VOC:\n- 영향:\n"
                "# 개선과제 및 대응방안\n"
                "1. 과제명 — 우선순위, 방법, 기대효과\n\n"
                f"{digest}"
            ),
            expected_output="산업군별 주요 이슈와 개선과제·대응방안 마크다운",
            agent=issue_agent,
        )
        result = run_crew([issue_agent], [issue_task], logger, "Issue")
        text = ""
        try:
            if issue_task.output:
                text = str(issue_task.output.raw or issue_task.output)
        except Exception:
            text = ""
        if not text:
            text = raw_text(result)
        return (text or "").strip()
    finally:
        issue_task = None
        issue_agent = None
        crew = None
        llm = None


def do_report(logger: Callable[[str], None]) -> dict[str, Any]:
    df = require_df()
    with STORE_LOCK:
        STORE["issue_analysis"] = ""
        STORE["report_path"] = None
        STORE["report_missing_images"] = []

    missing_images: list[str] = []
    logger("통계와 워드클라우드를 이 컴퓨터에서 준비합니다.")
    try:
        _, _, industry_png = make_ratio_bar(df, "산업군")
        if industry_png is None:
            missing_images.append("산업군 비율 그림")
            logger("산업군 비율 PNG를 저장하지 못했습니다.")
    except Exception as exc:
        missing_images.append("산업군 비율 그림")
        logger(f"산업군 통계 실패: {exc}")
        with STORE_LOCK:
            STORE["stats"]["산업군"] = ratio_table(df, "산업군")
    try:
        _, _, field_png = make_ratio_bar(df, "분야")
        if field_png is None:
            missing_images.append("분야 비율 그림")
            logger("분야 비율 PNG를 저장하지 못했습니다.")
    except Exception as exc:
        missing_images.append("분야 비율 그림")
        logger(f"분야 통계 실패: {exc}")
        with STORE_LOCK:
            STORE["stats"]["분야"] = ratio_table(df, "분야")
    try:
        wc_path = STORE.get("wordcloud_path")
        if wc_path is None or not Path(wc_path).exists():
            make_wordcloud(df)
    except Exception as exc:
        missing_images.append("워드클라우드")
        logger(f"워드클라우드 실패: {exc}")
        with STORE_LOCK:
            STORE["wordcloud_path"] = None

    digest = build_voc_digest(df)
    logger("통계 요약 작성 완료. 이슈 문장만 OpenAI를 사용합니다.")

    issue_ok = False
    api_key = resolve_api_key()
    if not api_key:
        logger("API 키가 없어 이슈 문장을 만들지 않습니다.")
        issue_text = (
            "# 이슈 분석을 만들지 못했습니다\n"
            "OpenAI API 키가 없어 이번 작업에서 산업군별 이슈 문장을 생성하지 않았습니다. "
            "통계와 그림은 이 컴퓨터에서 계산한 값입니다.\n"
        )
    else:
        try:
            issue_text = run_issue_analysis(digest, api_key, logger)
            if issue_text:
                issue_ok = True
            else:
                issue_text = (
                    "# 이슈 분석을 만들지 못했습니다\n"
                    "에이전트가 이번 작업에서 이슈 문장을 반환하지 않았습니다.\n"
                )
        except Exception as exc:
            logger(f"이슈 생성 실패: {exc}")
            issue_text = (
                "# 이슈 분석을 만들지 못했습니다\n"
                f"{friendly_error(exc)}\n"
            )

    logger("이번 작업의 이슈 내용으로 워드 파일을 저장합니다.")
    report_path = generate_word_report(issue_text, missing_images=missing_images)
    missing_final = list(STORE.get("report_missing_images") or missing_images)
    logger(f"워드 보고서 저장: {Path(report_path).name}")
    return {
        "path": str(report_path),
        "preview": issue_text,
        "issue_ok": issue_ok,
        "missing_images": missing_final,
    }


def stream_logs(work: JobFn, api_key: str | None = None):
    log_q: queue.Queue = queue.Queue()
    box: dict[str, Any] = {}

    def logger(msg: str) -> None:
        log_q.put(("log", f"[{now_ts()}] {sanitize_log(msg)}"))

    def worker() -> None:
        set_runtime_api_key(api_key)
        try:
            with JOB_LOCK:
                box["out"] = work(logger)
            log_q.put(("ok", None))
        except Exception as exc:
            box["err"] = exc
            log_q.put(("err", exc))
        finally:
            set_runtime_api_key(None)
            setattr(_RUNTIME, "llm", None)

    threading.Thread(target=worker, daemon=True).start()
    lines = [f"[{now_ts()}] 작업을 시작합니다."]
    yield "run", "\n".join(lines), None
    while True:
        try:
            kind, payload = log_q.get(timeout=0.25)
        except queue.Empty:
            yield "run", "\n".join(lines), None
            continue
        if kind == "log":
            lines.append(payload)
            yield "run", "\n".join(lines), None
        elif kind == "err":
            lines.append(f"[{now_ts()}] 처리에 실패했습니다.")
            yield "err", "\n".join(lines), payload
            return
        else:
            lines.append(f"[{now_ts()}] 결과 표시를 마쳤습니다.")
            yield "ok", "\n".join(lines), box.get("out")
            return


APP_CSS = """
.gradio-container { max-width: 1120px !important; }
.app-header { background:#fff; border:1px solid #D0D5DD; border-radius:12px; padding:16px 18px; margin-bottom:14px; }
.app-kicker { font-size:12px; color:#667085; margin-bottom:4px; }
.app-title { font-size:22px; font-weight:700; color:#1F3A5F; line-height:1.3; }
.app-header-meta { display:flex; gap:16px; flex-wrap:wrap; margin-top:12px; }
.meta-block { min-width:160px; }
.meta-k { font-size:11px; color:#667085; margin-bottom:2px; }
.meta-v { font-size:15px; font-weight:700; color:#1D2939; }
.meta-s { font-size:12px; color:#667085; margin-top:2px; }
.badge { display:inline-block; padding:2px 8px; border-radius:999px; font-size:12px; font-weight:700; }
.badge-idle { background:#F2F4F7; color:#344054; }
.badge-busy { background:#FEF0C7; color:#93370D; }
.badge-ready { background:#D1FADF; color:#05603A; }
.badge-fail { background:#FEE4E2; color:#B42318; }
.cards { display:grid; grid-template-columns:repeat(auto-fit,minmax(180px,1fr)); gap:12px; margin:8px 0 16px; }
.card { background:#fff; border:1px solid #EAECF0; border-radius:10px; padding:14px 16px; }
.card .k { font-size:12px; color:#667085; }
.card .v { font-size:20px; font-weight:700; color:#1F3A5F; margin:6px 0 4px; line-height:1.25; }
.card .s { font-size:12px; color:#667085; }
.result-head h2 { font-size:18px; font-weight:700; color:#1D2939; margin:0 0 4px; }
.result-sub { font-size:13px; color:#475467; margin:0 0 10px; }
.notice { border-radius:10px; padding:12px 14px; margin:8px 0 14px; font-size:14px; line-height:1.5; }
.notice-t { font-weight:700; margin-bottom:4px; }
.notice-info { background:#F8FAFC; border:1px solid #EAECF0; color:#344054; }
.notice-busy { background:#FFFAEB; border:1px solid #FEDF89; color:#93370D; }
.notice-fail { background:#FEF3F2; border:1px solid #FECDCA; color:#B42318; }
button.source-chip { background:#fff !important; color:#1F3A5F !important; border:1px solid #D0D5DD !important; font-weight:600 !important; }
button.source-chip:hover { background:#F2F4F7 !important; color:#1F3A5F !important; }
button.source-chip.is-selected { background:#1F3A5F !important; color:#fff !important; border:2px solid #163056 !important; }
button[role="tab"][aria-selected="true"] { font-weight:700 !important; color:#1F3A5F !important; }
.kw-row { display:flex; flex-wrap:wrap; gap:8px; margin:8px 0; }
.kw { background:#F2F4F7; border-radius:999px; padding:4px 10px; font-size:13px; color:#1D2939; }
.kw em { font-style:normal; color:#1F3A5F; font-weight:700; }
.muted { color:#667085; font-size:13px; }
.log-box textarea { font-family:Consolas,'Malgun Gothic',monospace; font-size:12px; }
"""


def build_ui() -> gr.Blocks:
    initial = empty_state()
    with gr.Blocks(title="고객사 VOC 분석 Agent") as demo:
        app_state = gr.State(initial)
        header = gr.HTML(render_header(initial))
        with gr.Accordion("OpenAI API 키 (이슈 분석·에이전트, 선택)", open=False):
            api_key_in = gr.Textbox(
                label="API 키",
                type="password",
                value="",
                placeholder="비우면 환경 변수 OPENAI_API_KEY를 사용합니다. 앱은 키를 저장하지 않습니다.",
                lines=1,
                max_lines=1,
            )
            gr.Markdown(
                "통계·워드클라우드·DOCX 저장은 키 없이 이 컴퓨터에서 처리합니다. "
                "이슈 문장을 만들 때만 통계 요약과 대표 불만 문구가 OpenAI로 전송됩니다. "
                "불만 문구에 이름·연락처 등 개인정보가 있으면 함께 전송될 수 있습니다. "
                "화면에서 입력한 키는 작업이 끝나면 입력란에서 지웁니다. "
                "환경 변수나 `.env`는 앱이 수정하지 않습니다."
            )
        with gr.Tabs():
            with gr.Tab("데이터 불러오기") as tab_upload:
                upload_notice = gr.HTML(INFO_UPLOAD)
                file_in = gr.File(label="VOC CSV 선택", file_types=[".csv"], type="filepath", file_count="single")
                with gr.Row():
                    btn_upload = gr.Button("이 파일 적용", variant="primary")
                    btn_sample = gr.Button("실습 샘플 불러오기", elem_classes=["source-chip"])
                upload_cards = gr.HTML("")
                df_preview = gr.Dataframe(
                    label="VOC 미리보기 (불만은 앞부분만 표시)",
                    value=EMPTY_PREVIEW,
                    wrap=True,
                    max_height=320,
                    interactive=False,
                    show_search="search",
                )
                with gr.Accordion("불만 원문 전체 보기", open=False):
                    df_full = gr.Dataframe(label="전체 VOC", value=EMPTY_FULL, wrap=True, max_height=360, interactive=False)
                with gr.Accordion("처리 기록", open=False):
                    log_upload = gr.Textbox(label="기록", lines=8, autoscroll=True, elem_classes=["log-box"])

            with gr.Tab("통계분석") as tab_stats:
                stats_notice = gr.HTML(INFO_STATS)
                col_dd = gr.Dropdown(label="비율을 볼 기준", choices=STAT_COLUMNS, value="산업군")
                btn_stats = gr.Button("비율 계산", variant="primary")
                stats_cards = gr.HTML("")
                plot_view = gr.Plot(label="구분별 VOC 비율(%)")
                stats_table = gr.Dataframe(label="구분별 건수와 비율", value=EMPTY_STATS, wrap=True, interactive=False, max_height=280)
                with gr.Accordion("처리 기록", open=False):
                    log_stats = gr.Textbox(label="기록", lines=8, autoscroll=True, elem_classes=["log-box"])

            with gr.Tab("워드클라우드") as tab_wc:
                wc_notice = gr.HTML(INFO_WC)
                btn_wc = gr.Button("키워드 그림 만들기", variant="primary")
                wc_cards = gr.HTML("")
                wc_image = gr.Image(label="불만 키워드 그림", type="filepath", interactive=False, format="png", height=420)
                with gr.Accordion("처리 기록", open=False):
                    log_wc = gr.Textbox(label="기록", lines=8, autoscroll=True, elem_classes=["log-box"])

            with gr.Tab("보고서생성") as tab_report:
                report_notice = gr.HTML(INFO_REPORT)
                btn_report = gr.Button("보고서 만들기", variant="primary")
                report_file = gr.File(label="워드 파일", file_types=[".docx"], interactive=False)
                issue_md = gr.Markdown("")
                with gr.Accordion("처리 기록", open=False):
                    log_report = gr.Textbox(label="기록", lines=8, autoscroll=True, elem_classes=["log-box"])

        def on_tab(state: dict[str, Any], title: str):
            state = dict(state)
            state["view_title"] = title
            return state, render_header(state)

        def cleared_analysis():
            return (
                INFO_STATS, "", None, EMPTY_STATS,
                INFO_WC, "", None,
                INFO_REPORT, None, "",
            )

        def pack_upload(state, notice, logs, preview, full, cards, busy, clear_file=True):
            return (
                state,
                render_header(state),
                notice,
                logs,
                preview,
                full,
                cards,
                *control_updates(busy, state, clear_file=clear_file),
                *cleared_analysis(),
            )

        def ui_upload(file_obj, state: dict[str, Any]):
            state = dict(state)
            keep_preview, keep_full, keep_cards = current_upload_views(state)
            if file_obj is None:
                state["status"] = "failed"
                state["status_text"] = "실패"
                keep = f"현재는 {state.get('source_label')} 데이터를 그대로 둡니다." if state.get("has_data") else "아직 적용된 데이터가 없습니다."
                notice = render_notice("fail", "적용할 CSV가 없습니다", f"파일을 선택한 뒤 [이 파일 적용]을 누르세요. {keep}")
                yield (
                    state, render_header(state), notice, f"[{now_ts()}] CSV가 선택되지 않았습니다.",
                    keep_preview, keep_full, keep_cards, *control_updates(False, state, True),
                    gr.skip(), gr.skip(), gr.skip(), gr.skip(),
                    gr.skip(), gr.skip(), gr.skip(),
                    gr.skip(), gr.skip(), gr.skip(),
                )
                return
            try:
                dest, name = persist_user_csv(file_obj)
            except Exception as exc:
                state["status"] = "failed"
                state["status_text"] = "실패"
                keep = f"새 파일은 적용되지 않았습니다. 현재는 {state.get('source_label')} 데이터를 표시합니다." if state.get("has_data") else "적용된 데이터가 없습니다."
                notice = render_notice("fail", "파일을 읽지 못했습니다", f"{friendly_error(exc)} {keep}")
                yield (
                    state, render_header(state), notice, f"[{now_ts()}] 파일 확보 실패",
                    keep_preview, keep_full, keep_cards, *control_updates(False, state, True),
                    gr.skip(), gr.skip(), gr.skip(), gr.skip(),
                    gr.skip(), gr.skip(), gr.skip(),
                    gr.skip(), gr.skip(), gr.skip(),
                )
                return

            state["busy"] = True
            state["status"] = "uploading"
            state["status_text"] = "불러오는 중"
            yield (
                state, render_header(state),
                render_notice("busy", "파일을 적용하는 중", f"{name}을 읽고 있습니다."),
                f"[{now_ts()}] {name} 내용을 확보한 뒤 적용합니다.",
                EMPTY_PREVIEW, EMPTY_FULL, "", *control_updates(True, state, True),
                gr.skip(), gr.skip(), gr.skip(), gr.skip(),
                gr.skip(), gr.skip(), gr.skip(),
                gr.skip(), gr.skip(), gr.skip(),
            )

            result = None
            logs = ""
            err = None
            for phase, logs, payload in stream_logs(lambda lg: do_upload(dest, "file", name, lg)):
                if phase == "run":
                    yield (
                        state, render_header(state),
                        render_notice("busy", "파일을 적용하는 중", "열과 건수를 확인하고 있습니다."),
                        logs, EMPTY_PREVIEW, EMPTY_FULL, "", *control_updates(True, state, True),
                        gr.skip(), gr.skip(), gr.skip(), gr.skip(),
                        gr.skip(), gr.skip(), gr.skip(),
                        gr.skip(), gr.skip(), gr.skip(),
                    )
                elif phase == "err":
                    err = payload
                else:
                    result = payload
            if err is not None:
                state["busy"] = False
                state["status"] = "failed"
                state["status_text"] = "실패"
                keep = f"새 파일은 적용되지 않았습니다. 현재는 {state.get('source_label')} 데이터를 표시합니다." if state.get("has_data") else "적용된 데이터가 없습니다."
                notice = render_notice("fail", "파일 적용에 실패했습니다", f"{friendly_error(err)} {keep}")
                yield (
                    state, render_header(state), notice, logs,
                    keep_preview, keep_full, keep_cards, *control_updates(False, state, True),
                    gr.skip(), gr.skip(), gr.skip(), gr.skip(),
                    gr.skip(), gr.skip(), gr.skip(),
                    gr.skip(), gr.skip(), gr.skip(),
                )
                return
            new_state = result["state"]
            yield (
                new_state, render_header(new_state),
                render_notice("info", "이 파일로 전환했습니다", "이전 통계·워드클라우드·보고서는 새 데이터와 섞이지 않도록 비웠습니다."),
                logs, result["preview"], result["full"], result["cards"],
                *control_updates(False, new_state, True),
                *cleared_analysis(),
            )

        def ui_sample(state: dict[str, Any]):
            state = dict(state)
            keep_preview, keep_full, keep_cards = current_upload_views(state)
            if not SAMPLE_CSV.exists():
                state["status"] = "failed"
                state["status_text"] = "실패"
                keep = f"샘플은 적용되지 않았습니다. 현재는 {state.get('source_label')} 데이터를 표시합니다." if state.get("has_data") else "적용된 데이터가 없습니다."
                notice = render_notice("fail", "실습 샘플 파일이 없습니다", keep)
                yield (
                    state, render_header(state), notice, f"[{now_ts()}] 샘플 없음",
                    keep_preview, keep_full, keep_cards, *control_updates(False, state, True),
                    gr.skip(), gr.skip(), gr.skip(), gr.skip(),
                    gr.skip(), gr.skip(), gr.skip(),
                    gr.skip(), gr.skip(), gr.skip(),
                )
                return

            pending = dict(state)
            pending["busy"] = True
            pending["status"] = "uploading"
            pending["status_text"] = "불러오는 중"
            yield (
                pending, render_header(pending),
                render_notice("busy", "실습 샘플을 적용하는 중", "선택 표시는 불러오기가 성공한 뒤에만 켭니다."),
                f"[{now_ts()}] 실습 샘플을 읽습니다.",
                EMPTY_PREVIEW, EMPTY_FULL, "", *control_updates(True, {**state, "source_kind": None}, True),
                gr.skip(), gr.skip(), gr.skip(), gr.skip(),
                gr.skip(), gr.skip(), gr.skip(),
                gr.skip(), gr.skip(), gr.skip(),
            )
            result = None
            logs = ""
            err = None
            for phase, logs, payload in stream_logs(lambda lg: do_upload(SAMPLE_CSV, "sample", SAMPLE_CSV.name, lg)):
                if phase == "run":
                    yield (
                        pending, render_header(pending),
                        render_notice("busy", "실습 샘플을 적용하는 중", "선택 표시는 불러오기가 성공한 뒤에만 켭니다."),
                        logs, EMPTY_PREVIEW, EMPTY_FULL, "", *control_updates(True, {**state, "source_kind": None}, True),
                        gr.skip(), gr.skip(), gr.skip(), gr.skip(),
                        gr.skip(), gr.skip(), gr.skip(),
                        gr.skip(), gr.skip(), gr.skip(),
                    )
                elif phase == "err":
                    err = payload
                else:
                    result = payload
            if err is not None:
                state["status"] = "failed"
                state["status_text"] = "실패"
                keep = f"샘플은 적용되지 않았습니다. 현재는 {state.get('source_label')} 데이터를 표시합니다." if state.get("has_data") else "적용된 데이터가 없습니다."
                notice = render_notice("fail", "실습 샘플을 적용하지 못했습니다", f"{friendly_error(err)} {keep}")
                yield (
                    state, render_header(state), notice, logs,
                    keep_preview, keep_full, keep_cards, *control_updates(False, state, True),
                    gr.skip(), gr.skip(), gr.skip(), gr.skip(),
                    gr.skip(), gr.skip(), gr.skip(),
                    gr.skip(), gr.skip(), gr.skip(),
                )
                return
            new_state = result["state"]
            yield (
                new_state, render_header(new_state),
                render_notice("info", "실습 샘플이 적용되었습니다", "이전 분석 결과는 비웠습니다. 선택됨 표시는 현재 화면에 적용된 데이터와 같습니다."),
                logs, result["preview"], result["full"], result["cards"],
                *control_updates(False, new_state, True),
                *cleared_analysis(),
            )

        def ui_stats(column, state: dict[str, Any]):
            state = dict(state)
            if not state.get("has_data"):
                yield (
                    state, render_header(state),
                    render_notice("info", "표시할 데이터가 없습니다", "데이터 불러오기에서 CSV 또는 실습 샘플을 먼저 적용하세요."),
                    "", gr.skip(), gr.skip(), gr.skip(),
                    *control_updates(False, state, False),
                )
                return
            state["busy"] = True
            state["status"] = "analyzing"
            state["status_text"] = "분석 중"
            state["view_title"] = "통계분석"
            yield (
                state, render_header(state),
                render_notice("busy", "비율을 계산하는 중", f"{state.get('source_label')} · 기준 {column}. 진행률 숫자는 표시하지 않습니다."),
                f"[{now_ts()}] {column} 비율 계산을 시작합니다.",
                gr.skip(), gr.skip(), gr.skip(), *control_updates(True, state, False),
            )
            result = None
            logs = ""
            err = None
            try:
                for phase, logs, payload in stream_logs(lambda lg: do_stats(column, lg)):
                    if phase == "run":
                        yield (
                            state, render_header(state),
                            render_notice("busy", "비율을 계산하는 중", f"현재 작업: {column}별 VOC 건수와 비율."),
                            logs, gr.skip(), gr.skip(), gr.skip(), *control_updates(True, state, False),
                        )
                    elif phase == "err":
                        err = payload
                    else:
                        result = payload
            except Exception as exc:
                err = exc
            if err is not None:
                state["busy"] = False
                state["status"] = "failed"
                state["status_text"] = "실패"
                yield (
                    state, render_header(state),
                    render_notice("fail", "비율 계산에 실패했습니다", f"{friendly_error(err)} 기준을 확인한 뒤 다시 실행하세요."),
                    logs, gr.skip(), gr.skip(), gr.skip(), *control_updates(False, state, False),
                )
                return
            fig, table = result
            if table is None or table.empty:
                state["busy"] = False
                state["status"] = "done"
                state["status_text"] = "분석 완료"
                yield (
                    state, render_header(state),
                    render_notice("info", "집계할 구분이 없습니다", "값이 0인 항목이 숨겨진 것이 아니라, 해당 열에 표시할 값이 없습니다."),
                    logs, "", None, EMPTY_STATS, *control_updates(False, state, False),
                )
                return
            state["busy"] = False
            state["status"] = "done"
            state["status_text"] = "분석 완료"
            state["stats_ready"] = True
            state["stats_column"] = column
            yield (
                state, render_header(state), "", logs,
                render_stats_cards(table, column, state), fig, display_stats_table(table),
                *control_updates(False, state, False),
            )

        def ui_wordcloud(state: dict[str, Any]):
            state = dict(state)
            if not state.get("has_data"):
                yield (
                    state, render_header(state),
                    render_notice("info", "표시할 데이터가 없습니다", "데이터 불러오기에서 CSV 또는 실습 샘플을 먼저 적용하세요."),
                    "", gr.skip(), gr.skip(), *control_updates(False, state, False),
                )
                return
            state["busy"] = True
            state["status"] = "analyzing"
            state["status_text"] = "분석 중"
            state["view_title"] = "워드클라우드"
            yield (
                state, render_header(state),
                render_notice("busy", "키워드 그림을 만드는 중", f"{state.get('source_label')}의 불만 열을 읽고 있습니다."),
                f"[{now_ts()}] 워드클라우드 작업을 시작합니다.",
                gr.skip(), gr.skip(), *control_updates(True, state, False),
            )
            result = None
            logs = ""
            err = None
            try:
                for phase, logs, payload in stream_logs(lambda lg: do_wordcloud(lg)):
                    if phase == "run":
                        yield (
                            state, render_header(state),
                            render_notice("busy", "키워드 그림을 만드는 중", "한글 폰트로 빈도를 그리는 중입니다."),
                            logs, gr.skip(), gr.skip(), *control_updates(True, state, False),
                        )
                    elif phase == "err":
                        err = payload
                    else:
                        result = payload
            except Exception as exc:
                err = exc
            if err is not None:
                state["busy"] = False
                state["status"] = "failed"
                state["status_text"] = "실패"
                yield (
                    state, render_header(state),
                    render_notice("fail", "워드클라우드를 만들지 못했습니다", f"{friendly_error(err)} 다시 실행해 주세요."),
                    logs, gr.skip(), gr.skip(), *control_updates(False, state, False),
                )
                return
            path = result[0] if isinstance(result, tuple) else result
            state["busy"] = False
            state["status"] = "done"
            state["status_text"] = "분석 완료"
            state["wc_ready"] = True
            yield (
                state, render_header(state), "", logs, render_wc_cards(state), path,
                *control_updates(False, state, False),
            )

        def ui_report(state: dict[str, Any], api_key: str = ""):
            state = dict(state)
            key_keep = gr.skip()
            key_clear = gr.update(value="")
            if not state.get("has_data"):
                yield (
                    state, render_header(state),
                    render_notice("info", "표시할 데이터가 없습니다", "데이터 불러오기에서 CSV 또는 실습 샘플을 먼저 적용하세요."),
                    "", gr.skip(), gr.skip(), *control_updates(False, state, False), key_keep,
                )
                return
            state["busy"] = True
            state["status"] = "analyzing"
            state["status_text"] = "분석 중"
            state["view_title"] = "보고서생성"
            yield (
                state, render_header(state),
                render_notice("busy", "보고서를 만드는 중", f"{state.get('source_label')} 기준으로 로컬 통계 후 이슈 문장을 요청합니다."),
                f"[{now_ts()}] 보고서 작업을 시작합니다.",
                gr.skip(), gr.skip(), *control_updates(True, state, False), key_keep,
            )
            result = None
            logs = ""
            err = None
            try:
                for phase, logs, payload in stream_logs(lambda lg: do_report(lg), api_key=api_key):
                    if phase == "run":
                        yield (
                            state, render_header(state),
                            render_notice("busy", "보고서를 만드는 중", "이 단계는 시간이 걸릴 수 있습니다."),
                            logs, gr.skip(), gr.skip(), *control_updates(True, state, False), key_keep,
                        )
                    elif phase == "err":
                        err = payload
                    else:
                        result = payload
            except Exception as exc:
                err = exc
            if err is not None:
                state["busy"] = False
                state["status"] = "failed"
                state["status_text"] = "실패"
                yield (
                    state, render_header(state),
                    render_notice("fail", "보고서를 만들지 못했습니다", f"{friendly_error(err)} 잠시 후 다시 실행해 주세요."),
                    logs, gr.skip(), gr.skip(), *control_updates(False, state, False), key_clear,
                )
                return
            path = result["path"] if isinstance(result, dict) else result[0]
            preview = result["preview"] if isinstance(result, dict) else result[1]
            issue_ok = result.get("issue_ok", False) if isinstance(result, dict) else True
            missing = result.get("missing_images") or [] if isinstance(result, dict) else []
            state["busy"] = False
            state["status"] = "done"
            state["report_ready"] = True
            notes = []
            if not issue_ok:
                notes.append("이슈 문장은 이번 작업에서 만들지 못했습니다.")
                state["status_text"] = "보고서 저장됨 · 이슈 없음"
            elif missing:
                state["status_text"] = "보고서 저장됨 · 일부 그림 없음"
            else:
                state["status_text"] = "보고서 저장됨"
            if missing:
                notes.append("넣지 못한 그림이 있어 완전한 보고서는 아닙니다: " + ", ".join(str(m) for m in missing))
            if not notes:
                notes.append("아래 미리보기와 다운로드 파일은 이번 작업의 같은 이슈 내용입니다.")
            report_title = "VOC 분석 보고서 (그림 누락 · 불완전)" if missing else "VOC 분석 보고서"
            head = (
                f'<div class="result-head"><h2>{report_title}</h2>'
                f'<p class="result-sub">분석 대상: {esc(state.get("source_label"))} · {esc(" ".join(notes))} '
                f"자료에 없는 일정·담당자는 단정하지 않습니다.</p></div>"
            )
            yield (
                state, render_header(state), head, logs, path, preview,
                *control_updates(False, state, False), key_clear,
            )

        upload_outs = [
            app_state, header, upload_notice, log_upload, df_preview, df_full, upload_cards,
            file_in, btn_upload, btn_sample, btn_stats, btn_wc, btn_report,
            stats_notice, stats_cards, plot_view, stats_table,
            wc_notice, wc_cards, wc_image,
            report_notice, report_file, issue_md,
        ]
        stats_outs = [
            app_state, header, stats_notice, log_stats, stats_cards, plot_view, stats_table,
            file_in, btn_upload, btn_sample, btn_stats, btn_wc, btn_report,
        ]
        wc_outs = [
            app_state, header, wc_notice, log_wc, wc_cards, wc_image,
            file_in, btn_upload, btn_sample, btn_stats, btn_wc, btn_report,
        ]
        report_outs = [
            app_state, header, report_notice, log_report, report_file, issue_md,
            file_in, btn_upload, btn_sample, btn_stats, btn_wc, btn_report,
            api_key_in,
        ]

        event_kw = dict(show_progress="minimal", concurrency_id="voc-app", concurrency_limit=1)
        tab_kw = dict(show_progress="hidden")
        tab_upload.select(lambda s: on_tab(s, "데이터 불러오기"), inputs=[app_state], outputs=[app_state, header], **tab_kw)
        tab_stats.select(lambda s: on_tab(s, "통계분석"), inputs=[app_state], outputs=[app_state, header], **tab_kw)
        tab_wc.select(lambda s: on_tab(s, "워드클라우드"), inputs=[app_state], outputs=[app_state, header], **tab_kw)
        tab_report.select(lambda s: on_tab(s, "보고서생성"), inputs=[app_state], outputs=[app_state, header], **tab_kw)
        btn_upload.click(ui_upload, inputs=[file_in, app_state], outputs=upload_outs, **event_kw)
        btn_sample.click(ui_sample, inputs=[app_state], outputs=upload_outs, **event_kw)
        btn_stats.click(ui_stats, inputs=[col_dd, app_state], outputs=stats_outs, **event_kw)
        btn_wc.click(ui_wordcloud, inputs=[app_state], outputs=wc_outs, **event_kw)
        btn_report.click(ui_report, inputs=[app_state, api_key_in], outputs=report_outs, **event_kw)
    return configure_file_access(demo)


def configure_file_access(demo: gr.Blocks) -> gr.Blocks:
    demo.allowed_paths = [str(EXPORT_DIR.resolve())]
    demo.blocked_paths = default_blocked_paths()
    install_download_guard()
    return demo


def main() -> None:
    ok, msg = python_support_message()
    if not ok:
        rec = ".".join(str(x) for x in RECOMMENDED_PYTHON)
        kr = (
            f"이 앱은 Python 3.11, 3.12, 3.13에서만 설치·실행할 수 있습니다. "
            f"현재는 Python {sys.version_info.major}.{sys.version_info.minor}.{sys.version_info.micro}입니다. "
            f"고정된 CrewAI 1.15.1은 Python 3.14를 허용하지 않고, pandas 3.0은 3.11 이상이 필요합니다. "
            f"Python {rec}을 권장합니다."
        )
        print(kr, file=sys.stderr)
        print(msg, file=sys.stderr)
        raise SystemExit(1)
    if FONT_WARNING:
        print("경고: " + FONT_WARNING, file=sys.stderr)
    if port_in_use(APP_PORT):
        print(
            f"포트 {APP_PORT}이 이미 사용 중입니다. "
            f"실행 중인 앱을 종료하거나 http://127.0.0.1:{APP_PORT} 으로 접속하세요.",
            file=sys.stderr,
        )
        raise SystemExit(1)
    demo = configure_file_access(build_ui())
    demo.launch(
        server_name="127.0.0.1",
        server_port=APP_PORT,
        inbrowser=True,
        theme=gr.themes.Soft(font=[FONT_FAMILY, "sans-serif"]),
        css=APP_CSS,
        footer_links=["gradio"],
        allowed_paths=[str(EXPORT_DIR.resolve())],
        blocked_paths=default_blocked_paths(),
        show_error=False,
    )


if __name__ == "__main__":
    main()
